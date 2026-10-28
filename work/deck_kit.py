# -*- coding: utf-8 -*-
"""古风汇报 PPT 版式工具：纸色底、墨字、印章红点缀。"""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

# ---------- 配色 ----------
PAPER = RGBColor(0xEF, 0xE7, 0xD5)
PAPER_TOP = RGBColor(0xF7, 0xF2, 0xE6)
PAPER_DEEP = RGBColor(0xE4, 0xD8, 0xC0)
CARD = RGBColor(0xFB, 0xF7, 0xEE)
INK = RGBColor(0x1B, 0x19, 0x16)
INK_SOFT = RGBColor(0x45, 0x3D, 0x34)
MUTED = RGBColor(0x7C, 0x73, 0x64)
FAINT = RGBColor(0xA8, 0x9E, 0x8C)
SEAL = RGBColor(0x9C, 0x2F, 0x26)
SEAL_LT = RGBColor(0xB8, 0x44, 0x3A)
JADE = RGBColor(0x2F, 0x4A, 0x40)
JADE_LT = RGBColor(0x40, 0x69, 0x5A)
GOLD = RGBColor(0xB0, 0x8D, 0x57)
LINE = RGBColor(0xD5, 0xC8, 0xAE)
WATER = RGBColor(0xDE, 0xD2, 0xB8)

# ---------- 字体 ----------
KAI = "KaiTi"
SONG = "SimSun"
HEI = "Microsoft YaHei"

SW, SH = 13.333, 7.5

# 页序计数器：每新建一页自动 +1，供页脚页码使用，避免插入新页后手工改号。
_PAGE = {"n": 0}


def cur_page():
    return _PAGE["n"]


def E(v):
    return Inches(v)


def set_font(run, name):
    """同时设置拉丁、东亚、复杂文种字形，避免中文回退成默认字体。"""
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def gradient(shape, c1, c2, angle=90):
    fill = shape.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = c1
    stops[0].position = 0.0
    stops[1].color.rgb = c2
    stops[1].position = 1.0
    try:
        fill.gradient_angle = angle
    except Exception:
        pass


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE,
         dash=None, radius=None):
    sp = slide.shapes.add_shape(shape, E(x), E(y), E(w), E(h))
    no_shadow(sp)
    if radius is not None:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
        if dash is not None:
            sp.line.dash_style = dash
    sp.text_frame.word_wrap = True
    return sp


def hline(slide, x, y, w, color=LINE, lw=0.75):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x), E(y), E(x + w), E(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(lw)
    return ln


def vline(slide, x, y, h, color=LINE, lw=0.75):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x), E(y), E(x), E(y + h))
    ln.line.color.rgb = color
    ln.line.width = Pt(lw)
    return ln


def text(slide, x, y, w, h, items, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         wrap=True, vertical=None, margins=0.0):
    """items: list of dict(t=, s=, c=, b=, f=, al=, ls=, sb=, sa=, bullet=)"""
    box = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = E(margins)
    tf.vertical_anchor = anchor
    if vertical:
        bodyPr = tf._txBody.find(qn("a:bodyPr"))
        bodyPr.set("vert", vertical)
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = it.get("al", align)
        p.line_spacing = it.get("ls", 1.3)
        if it.get("sb"):
            p.space_before = Pt(it["sb"])
        if it.get("sa") is not None:
            p.space_after = Pt(it["sa"])
        else:
            p.space_after = Pt(0)
        content = it.get("t", "")
        if it.get("bullet"):
            content = it["bullet"] + " " + content
        run = p.add_run()
        run.text = content
        f = run.font
        f.size = Pt(it.get("s", 13))
        f.bold = it.get("b", False)
        f.color.rgb = it.get("c", INK)
        set_font(run, it.get("f", HEI))
    return box


def shape_text(shape, items, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margins=0.06):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = E(margins)
    tf.margin_top = tf.margin_bottom = E(margins * 0.5)
    tf.vertical_anchor = anchor
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = it.get("al", align)
        p.line_spacing = it.get("ls", 1.25)
        p.space_after = Pt(it.get("sa", 0))
        if it.get("sb"):
            p.space_before = Pt(it["sb"])
        run = p.add_run()
        run.text = it.get("t", "")
        f = run.font
        f.size = Pt(it.get("s", 12))
        f.bold = it.get("b", False)
        f.color.rgb = it.get("c", INK)
        set_font(run, it.get("f", HEI))
    return shape


# ---------- 页面骨架 ----------

def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = rect(slide, 0, 0, SW, SH, fill=PAPER)
    gradient(bg, PAPER_TOP, PAPER_DEEP, angle=45)
    _PAGE["n"] += 1
    return slide


def frame(slide, idx=None, total=None, section="", seal_text="論語"):
    """统一的页眉页脚与左侧装饰线。"""
    vline(slide, 0.62, 1.02, 5.72, color=RGBColor(0xDD, 0xD0, 0xB6), lw=2.0)
    if seal_text:
        sq = rect(slide, 0.70, 0.42, 0.36, 0.36, fill=SEAL)
        shape_text(sq, [{"t": seal_text[:1], "s": 12, "c": RGBColor(0xFF, 0xF8, 0xEE), "f": KAI}],
                   margins=0.0)
    if section:
        # 右上角章节标记：红色印章数字 + 章名，字号放大，便于远距离看清
        if "·" in section:
            num, name = [p.strip() for p in section.split("·", 1)]
        else:
            num, name = "", section
        if num:
            sq = rect(slide, 12.06, 0.36, 0.55, 0.55, fill=SEAL,
                      shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
            shape_text(sq, [{"t": num, "s": 21, "c": RGBColor(0xFF, 0xF8, 0xEE),
                             "f": KAI, "b": True}], margins=0.0)
            vline(slide, 8.10, 0.44, 0.40, color=GOLD, lw=1.25)
            text(slide, 8.28, 0.36, 3.62, 0.55,
                 [{"t": name, "s": 14.5, "c": INK_SOFT, "f": KAI, "b": True,
                   "al": PP_ALIGN.RIGHT}], anchor=MSO_ANCHOR.MIDDLE)
        else:
            text(slide, 8.28, 0.36, 3.62, 0.55,
                 [{"t": name, "s": 13.5, "c": GOLD, "f": HEI, "al": PP_ALIGN.RIGHT}],
                 anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 0.70, 6.98, 6.0, 0.3,
         [{"t": "孔子之境 · 智启论语 ｜ 人工智能+X 课程汇报 ｜ 深藏blue",
           "s": 8.5, "c": FAINT, "f": HEI}])
    if idx:
        text(slide, 11.2, 6.94, 1.43, 0.34,
             [{"t": "%02d / %02d" % (idx, total or idx), "s": 9.5, "c": MUTED, "f": HEI,
               "al": PP_ALIGN.RIGHT}])


def head(slide, title, sub=None, kicker=None):
    if kicker:
        text(slide, 1.10, 0.36, 8.0, 0.26,
             [{"t": kicker, "s": 10, "c": GOLD, "f": HEI}])
        ty = 0.60
    else:
        ty = 0.52
    text(slide, 1.10, ty, 9.6, 0.52,
         [{"t": title, "s": 26, "c": INK, "f": KAI, "b": True}])
    hline(slide, 1.10, 1.24, 11.5, color=RGBColor(0xE0, 0xD4, 0xBB), lw=1.0)
    if sub:
        text(slide, 1.10, 1.33, 11.4, 0.32,
             [{"t": sub, "s": 12.2, "c": MUTED, "f": HEI}])
        return 1.80
    return 1.52


def watermark(slide, ch, x=11.55, y=4.55, size=150):
    text(slide, x, y, 1.7, 2.6,
         [{"t": ch, "s": size, "c": WATER, "f": KAI, "al": PP_ALIGN.RIGHT}],
         vertical="eaVert")


def card(slide, x, y, w, h, title, lines, accent=SEAL, title_size=13, size=10.8,
         fill=CARD, tag=None):
    box = rect(slide, x, y, w, h, fill=fill, line=RGBColor(0xE2, 0xD7, 0xBF),
               lw=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(slide, x, y + 0.14, 0.055, h - 0.28, fill=accent)
    items = [{"t": title, "s": title_size, "c": INK, "f": KAI, "b": True, "sa": 6}]
    for ln in lines:
        items.append({"t": ln, "s": size, "c": INK_SOFT, "f": HEI, "ls": 1.34, "sa": 4})
    text(slide, x + 0.24, y + 0.16, w - 0.44, h - 0.3, items)
    if tag:
        tg = rect(slide, x + w - 0.72, y + 0.13, 0.58, 0.26, fill=accent,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
        shape_text(tg, [{"t": tag, "s": 8.5, "c": RGBColor(0xFF, 0xF8, 0xEE), "f": HEI, "b": True}],
                   margins=0.0)
    return box


def chip(slide, x, y, w, h, label, fill=None, color=INK_SOFT, size=10.5, border=None,
         font=HEI, bold=False):
    sp = rect(slide, x, y, w, h, fill=fill if fill is not None else RGBColor(0xF6, 0xF0, 0xE2),
              line=border if border is not None else RGBColor(0xE0, 0xD4, 0xBB),
              lw=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.28)
    shape_text(sp, [{"t": label, "s": size, "c": color, "f": font, "b": bold}], margins=0.05)
    return sp


def placeholder(slide, x, y, w, h, tag, title, hint, slot=None):
    """预留图片位：虚线框 + 注明需要放什么图。"""
    box = rect(slide, x, y, w, h, fill=RGBColor(0xF7, 0xF2, 0xE5), line=GOLD, lw=1.25,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03,
               dash=MSO_LINE_DASH_STYLE.DASH)
    tg = rect(slide, x + 0.14, y + 0.13, 1.00, 0.30, fill=SEAL,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    shape_text(tg, [{"t": tag, "s": 9.5, "c": RGBColor(0xFF, 0xF8, 0xEE), "f": HEI, "b": True}],
               margins=0.0)
    items = [{"t": title, "s": 14.5, "c": SEAL, "f": KAI, "b": True, "sa": 9, "al": PP_ALIGN.CENTER}]
    for ln in hint.split("\n"):
        items.append({"t": ln, "s": 11.3, "c": MUTED, "f": HEI, "ls": 1.48, "sa": 4,
                      "al": PP_ALIGN.CENTER})
    text(slide, x + 0.3, y + h / 2 - 1.05, w - 0.6, 2.1, items, anchor=MSO_ANCHOR.MIDDLE)
    if slot:
        text(slide, x + 0.16, y + h - 0.34, w - 0.32, 0.24,
             [{"t": slot, "s": 9.2, "c": FAINT, "f": HEI}])
    return box


def grid(slide, x, y, w, col_ratios, rows, row_h=0.52, head=True, size=10.6,
         head_size=11, head_fill=JADE, zebra=True, aligns=None, pad=0.14):
    """自绘表格：行底色 + 单元格文字，边框用细分隔线。"""
    total = sum(col_ratios)
    widths = [w * r / total for r in col_ratios]
    cy = y
    for ri, row in enumerate(rows):
        is_head = head and ri == 0
        h = row_h[ri] if isinstance(row_h, (list, tuple)) else row_h
        bg = head_fill if is_head else (RGBColor(0xF9, 0xF5, 0xEB) if (zebra and ri % 2 == 0)
                                        else RGBColor(0xF3, 0xEC, 0xDC))
        rect(slide, x, cy, w, h, fill=bg)
        cx = x
        for ci, cell in enumerate(row):
            al = PP_ALIGN.LEFT if aligns is None else aligns[ci]
            if is_head:
                al = PP_ALIGN.CENTER if aligns is None else aligns[ci]
            text(slide, cx + pad, cy + 0.06, widths[ci] - 2 * pad, h - 0.12,
                 [{"t": cell, "s": head_size if is_head else size,
                   "c": RGBColor(0xFF, 0xF8, 0xEE) if is_head else INK_SOFT,
                   "f": KAI if is_head else HEI, "b": is_head, "ls": 1.28,
                   "al": al}],
                 anchor=MSO_ANCHOR.MIDDLE)
            cx += widths[ci]
        hline(slide, x, cy + h, w, color=RGBColor(0xE3, 0xD8, 0xC0), lw=0.6)
        cy += h
    return cy


def ribbon(slide, x, y, w, h, items, fill=JADE, size=11.5, radius=0.06):
    sp = rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    shape_text(sp, items, align=PP_ALIGN.LEFT, margins=0.22)
    return sp


def flow(slide, y, steps, x0=0.95, x1=12.4, h=0.62, gap=0.30, fill=JADE, size=11.5):
    n = len(steps)
    bw = ((x1 - x0) - gap * (n - 1)) / n
    cx = x0
    for i, s in enumerate(steps):
        sp = rect(slide, cx, y, bw, h, fill=fill if i % 2 == 0 else JADE_LT,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
        shape_text(sp, [{"t": s, "s": size, "c": RGBColor(0xFF, 0xF8, 0xEE), "f": KAI, "b": True}])
        if i < n - 1:
            text(slide, cx + bw, y, gap, h,
                 [{"t": "→", "s": 15, "c": GOLD, "f": HEI, "al": PP_ALIGN.CENTER}],
                 anchor=MSO_ANCHOR.MIDDLE)
        cx += bw + gap
    return bw
